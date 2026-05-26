# =============================================================================
# VIS 이슈어 온보딩 트래커  ·  DealMe Internal
# =============================================================================
import streamlit as st
import anthropic
import json
import re
import hashlib
import base64
import copy
import os
from io import BytesIO
from datetime import datetime

import openpyxl
import gspread
from google.oauth2.service_account import Credentials
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics


# ── 다국어 지원 (Language Support) ─────────────────────────────────────────────
if "lang" not in st.session_state:
    st.session_state["lang"] = "ko"

_KO = {
    "login_title": "VIS 온보딩 트래커", "login_sub": "DealMe · 내부 전용",
    "login_pw": "비밀번호", "login_pw_ph": "비밀번호를 입력하세요",
    "login_btn": "로그인", "login_err": "비밀번호가 올바르지 않습니다.",
    "lang_toggle": "🌐 English",
    "settings": "⚙️ 설정", "api_key_set": "✅ API 키 설정됨",
    "api_key_warn": "AI 분석을 위해 API 키를 입력하세요",
    "gsheet_hdr": "📊 Google Sheets", "sheet_id_ph": "구글 시트 URL의 ID",
    "gsheet_ok": "✅ 구글 시트 연결됨", "save_btn": "💾 현재 상태 저장",
    "save_ok": "저장 완료! 변경 {n}건", "save_fail": "실패: {e}",
    "gsheet_warn": "Google Sheets 미연결",
    "bank_mgmt": "🏦 은행 관리", "add_bank_label": "은행 이름 추가",
    "add_bank_ph": "예: BIDV, VPBank", "add_btn": "➕ 추가",
    "add_warn_empty": "이름을 입력해주세요.",
    "add_warn_dup": "이미 등록되어 있습니다: {n}",
    "add_ok": "추가 완료: {n}",
    "del_bank_label": "은행 삭제", "del_select": "선택...",
    "del_btn": "🗑️ 삭제", "del_ok": "삭제 완료: {n}",
    "legend": "**범례** ✅완료 ❌미완료 🔄진행중 ❓확인중 ⚠️확인필요 🚨심각지연",
    "logout_btn": "🚪 로그아웃", "reset_btn": "🔄 초기화",
    "app_title": "💳 VIS 이슈어 온보딩 트래커",
    "app_sub": "신세계면세점 런칭 전 베트남 은행 온보딩 현황 | 마지막 업데이트: **{t}**",
    "kpi_total": "🏦 전체 은행", "kpi_ready": "✅ 런칭 준비완료",
    "kpi_inprog": "🔄 진행 중", "kpi_crit": "🚨 심각 지연",
    "kpi_tc": "📄 T&C 완료", "kpi_url": "🔗 URL 확보", "kpi_bin": "🔢 등록 BIN",
    "export_exp": "📤 보고서 내보내기",
    "export_xlsx": "📊 엑셀 다운로드 (.xlsx)", "export_pdf": "📄 PDF 보고서 (.pdf)",
    "tab_status": "📊 온보딩 현황", "tab_tc": "📄 T&C 현황", "tab_bin": "🔢 BIN 리스트",
    "tab1_hdr": "은행별 온보딩 현황",
    "fld_krw": "KRW/VND 플랜", "fld_usd": "USD 플랜", "fld_tc": "영어 T&C", "fld_bin": "BIN 상태",
    "edit_exp": "✏️ {b} 직접 편집", "save_edit_btn": "💾 저장", "save_edit_ok": "✅ {b} 저장 완료",
    "actions_hdr": "📌 현재 액션 아이템", "actions_empty": "등록된 액션 아이템이 없습니다.",
    "tab2_hdr": "카드사별 영문 T&C 현황",
    "tab2_cap": "이미지·이메일·엑셀 분석 시 T&C URL이 감지되면 자동 반영됩니다.",
    "tc_m_total": "전체 은행", "tc_m_done": "T&C 제출 완료", "tc_m_url": "URL 확보",
    "tc_col_bank": "**은행명**", "tc_col_stat": "**T&C 상태**",
    "tc_col_url": "**URL**", "tc_col_link": "**원문 링크**",
    "tc_view": "📄 T&C 원문 보기", "tc_no_url": "URL 미확보",
    "tc_edit_exp": "✏️ {b} T&C URL 수정", "tc_url_label": "T&C URL",
    "tc_save_btn": "저장", "tc_save_ok": "✅ {b} T&C URL 저장됨",
    "tab3_hdr": "🔢 은행별 BIN 리스트",
    "tab3_cap": "이메일·엑셀·이미지 분석 시 6~8자리 BIN 번호가 감지되면 자동으로 누적 추가됩니다. 수동 입력도 가능합니다.",
    "bin_m_with": "BIN 제출 은행", "bin_m_without": "BIN 미제출 은행", "bin_m_total": "전체 BIN 수",
    "bin_none": "아직 등록된 BIN이 없습니다.",
    "bin_exp": "📋 {b} BIN 전체 목록 보기 / 편집 ({n}개)",
    "bin_col_num": "**#**", "bin_col_bin": "**BIN 번호**", "bin_col_digits": "**자리수**",
    "bin_del_sel": "삭제할 BIN 선택", "bin_del_btn": "🗑️ 선택 BIN 삭제",
    "bin_del_ok": "✅ {b} 삭제됨", "bin_clear_btn": "⚠️ {b} BIN 전체 삭제",
    "bin_clear_ok": "✅ {b} BIN 목록 초기화됨",
    "bin_manual_exp": "➕ {b} BIN 수동 입력",
    "bin_manual_cap": "여러 개는 쉼표(,) 또는 줄바꿈으로 구분하세요. 예: 431167, 43116707, 456789",
    "bin_manual_label": "BIN 번호 입력",
    "bin_add_btn": "➕ 추가", "bin_add_ok": "✅ {added}개 추가됨 (전체 {total}개)",
    "bin_invalid": "⚠️ 유효하지 않은 항목 무시됨: {items} (6~8자리 숫자만 허용)",
    "input_hdr": "📥 새 데이터 입력 → AI 분석 → 자동 업데이트",
    "input_cap": "세 가지 방법 모두 동일한 템플릿에 반영됩니다. BIN 번호도 자동 추출·누적됩니다.",
    "inp_email": "📧 이메일 붙여넣기", "inp_excel": "📊 엑셀 업로드", "inp_img": "🖼️ 이미지 업로드",
    "email_cap": "아웃룩 이메일 스레드를 통째로 붙여넣으세요. AI가 최신 메시지 기준으로 상태·BIN·T&C URL 등을 자동 반영합니다.",
    "email_label": "이메일 본문", "email_ph": "이메일 전문을 그대로 붙여넣으세요...",
    "email_btn": "🤖 이메일 분석 후 업데이트",
    "need_key": "먼저 사이드바에서 API 키를 입력하세요.",
    "email_empty": "이메일 내용을 입력해주세요.",
    "dup_warn": "⚠️ 이미 처리한 내용입니다.",
    "spin_email": "AI가 이메일을 분석하는 중...",
    "update_ok": "✅ 현황판 업데이트 완료!", "new_banks": " 새 은행 추가: {banks}",
    "excel_cap": "은행이 보낸 온보딩 시트(.xlsx)를 업로드하세요. BIN 번호 목록도 자동 추출합니다.",
    "excel_upload": "엑셀 파일 (.xlsx)", "excel_uploaded": "📂 업로드됨: **{name}**",
    "excel_btn": "🤖 엑셀 분석 후 업데이트",
    "dup_file": "⚠️ 이미 처리한 파일입니다.",
    "spin_excel": "AI가 엑셀을 분석하는 중...",
    "excel_ok": "✅ 업데이트 완료! (제출 은행: {bank}, BIN: {cnt}개)",
    "img_cap": "이메일 스크린샷·표 캡처 등을 업로드하세요. AI가 BIN 번호와 T&C URL까지 이미지에서 직접 추출합니다.",
    "img_upload": "이미지 파일 (PNG, JPG, JPEG, WEBP)",
    "img_btn": "🤖 이미지 분석 후 업데이트",
    "dup_img": "⚠️ 이미 처리한 이미지입니다.",
    "spin_img": "AI가 이미지를 분석하는 중...",
    "irr_warn": "⚠️ {reason}",
    "irr_email": "온보딩과 관련 없는 내용입니다.",
    "irr_excel": "온보딩과 관련 없는 파일입니다.",
    "irr_img": "온보딩과 관련 없는 이미지입니다.",
    "err": "오류: {e}",
    "hist_exp": "🕓 업데이트 히스토리 ({n}개)",
    "hist_cap": "최대 10개 보관. 복원 버튼을 누르면 해당 시점으로 돌아갑니다.",
    "restore_btn": "↩ 복원", "restore_ok": "이전 상태로 복원했습니다.",
    "trend_hdr": "📈 날짜별 온보딩 추이",
    "trend_no_gs": "왼쪽 사이드바에서 Google Sheets를 연동하면 날짜별 추이를 볼 수 있습니다.",
    "trend_spin": "변경 이력 불러오는 중...",
    "trend_empty": "저장된 이력이 없습니다. 사이드바 '💾 현재 상태 저장'을 눌러주세요.",
    "trend_total": "**변경 이력** — 총 **{n}건**",
    "footer": "VIS Issuer Onboarding Tracker · Powered by Claude AI · DealMe Internal",
}

_EN = {
    "login_title": "VIS Onboarding Tracker", "login_sub": "DealMe · Internal Only",
    "login_pw": "Password", "login_pw_ph": "Enter password",
    "login_btn": "Login", "login_err": "Incorrect password.",
    "lang_toggle": "🌐 한국어",
    "settings": "⚙️ Settings", "api_key_set": "✅ API Key set",
    "api_key_warn": "Enter API key for AI analysis",
    "gsheet_hdr": "📊 Google Sheets", "sheet_id_ph": "Google Sheets URL ID",
    "gsheet_ok": "✅ Google Sheets connected", "save_btn": "💾 Save current state",
    "save_ok": "Saved! {n} changes", "save_fail": "Failed: {e}",
    "gsheet_warn": "Google Sheets not connected",
    "bank_mgmt": "🏦 Bank Management", "add_bank_label": "Add bank name",
    "add_bank_ph": "e.g. BIDV, VPBank", "add_btn": "➕ Add",
    "add_warn_empty": "Please enter a name.",
    "add_warn_dup": "Already registered: {n}",
    "add_ok": "Added: {n}",
    "del_bank_label": "Delete bank", "del_select": "Select...",
    "del_btn": "🗑️ Delete", "del_ok": "Deleted: {n}",
    "legend": "**Legend** ✅Done ❌Not done 🔄In progress ❓Checking ⚠️Review 🚨Critical",
    "logout_btn": "🚪 Logout", "reset_btn": "🔄 Reset",
    "app_title": "💳 VIS Issuer Onboarding Tracker",
    "app_sub": "Vietnam bank onboarding for Shinsegae duty-free | Last update: **{t}**",
    "kpi_total": "🏦 Total Banks", "kpi_ready": "✅ Launch Ready",
    "kpi_inprog": "🔄 In Progress", "kpi_crit": "🚨 Critical Delay",
    "kpi_tc": "📄 T&C Done", "kpi_url": "🔗 URL Secured", "kpi_bin": "🔢 Registered BINs",
    "export_exp": "📤 Export Report",
    "export_xlsx": "📊 Download Excel (.xlsx)", "export_pdf": "📄 PDF Report (.pdf)",
    "tab_status": "📊 Onboarding Status", "tab_tc": "📄 T&C Status", "tab_bin": "🔢 BIN List",
    "tab1_hdr": "Bank Onboarding Status",
    "fld_krw": "KRW/VND Plan", "fld_usd": "USD Plan", "fld_tc": "English T&C", "fld_bin": "BIN Status",
    "edit_exp": "✏️ Edit {b}", "save_edit_btn": "💾 Save", "save_edit_ok": "✅ {b} saved",
    "actions_hdr": "📌 Current Action Items", "actions_empty": "No action items registered.",
    "tab2_hdr": "T&C Status by Bank",
    "tab2_cap": "T&C URLs detected from image/email/excel analysis are automatically reflected.",
    "tc_m_total": "Total Banks", "tc_m_done": "T&C Submitted", "tc_m_url": "URL Secured",
    "tc_col_bank": "**Bank**", "tc_col_stat": "**T&C Status**",
    "tc_col_url": "**URL**", "tc_col_link": "**Document Link**",
    "tc_view": "📄 View T&C Document", "tc_no_url": "URL not available",
    "tc_edit_exp": "✏️ Edit {b} T&C URL", "tc_url_label": "T&C URL",
    "tc_save_btn": "Save", "tc_save_ok": "✅ {b} T&C URL saved",
    "tab3_hdr": "🔢 BIN List by Bank",
    "tab3_cap": "6-8 digit BIN numbers from email/excel/image are auto-accumulated. Manual entry also available.",
    "bin_m_with": "Banks with BIN", "bin_m_without": "Banks without BIN", "bin_m_total": "Total BINs",
    "bin_none": "No BINs registered yet.",
    "bin_exp": "📋 View/Edit {b} BIN List ({n} entries)",
    "bin_col_num": "**#**", "bin_col_bin": "**BIN Number**", "bin_col_digits": "**Digits**",
    "bin_del_sel": "Select BIN to delete", "bin_del_btn": "🗑️ Delete selected BIN",
    "bin_del_ok": "✅ {b} deleted", "bin_clear_btn": "⚠️ Clear all {b} BINs",
    "bin_clear_ok": "✅ {b} BIN list cleared",
    "bin_manual_exp": "➕ Manual BIN entry for {b}",
    "bin_manual_cap": "Separate with comma (,) or newline. e.g. 431167, 43116707, 456789",
    "bin_manual_label": "BIN Number Input",
    "bin_add_btn": "➕ Add", "bin_add_ok": "✅ {added} added (total {total})",
    "bin_invalid": "⚠️ Invalid entries ignored: {items} (only 6-8 digit numbers allowed)",
    "input_hdr": "📥 New Data Input → AI Analysis → Auto Update",
    "input_cap": "All three methods update the same tracker. BIN numbers are also auto-extracted.",
    "inp_email": "📧 Paste Email", "inp_excel": "📊 Upload Excel", "inp_img": "🖼️ Upload Image",
    "email_cap": "Paste the full Outlook email thread. AI reflects status, BIN, T&C URL from the latest message.",
    "email_label": "Email Body", "email_ph": "Paste the full email here...",
    "email_btn": "🤖 Analyze Email & Update",
    "need_key": "Please enter your API key in the sidebar first.",
    "email_empty": "Please enter email content.",
    "dup_warn": "⚠️ This content has already been processed.",
    "spin_email": "AI is analyzing the email...",
    "update_ok": "✅ Dashboard updated!", "new_banks": " New banks added: {banks}",
    "excel_cap": "Upload the onboarding sheet (.xlsx) sent by the bank. BIN numbers are also auto-extracted.",
    "excel_upload": "Excel file (.xlsx)", "excel_uploaded": "📂 Uploaded: **{name}**",
    "excel_btn": "🤖 Analyze Excel & Update",
    "dup_file": "⚠️ This file has already been processed.",
    "spin_excel": "AI is analyzing the Excel file...",
    "excel_ok": "✅ Update complete! (Bank: {bank}, BINs: {cnt})",
    "img_cap": "Upload email screenshots or table captures. AI extracts BIN numbers and T&C URLs from images.",
    "img_upload": "Image file (PNG, JPG, JPEG, WEBP)",
    "img_btn": "🤖 Analyze Image & Update",
    "dup_img": "⚠️ This image has already been processed.",
    "spin_img": "AI is analyzing the image...",
    "irr_warn": "⚠️ {reason}",
    "irr_email": "This content is not related to onboarding.",
    "irr_excel": "This file is not related to onboarding.",
    "irr_img": "This image is not related to onboarding.",
    "err": "Error: {e}",
    "hist_exp": "🕓 Update History ({n} entries)",
    "hist_cap": "Stores up to 10 snapshots. Click Restore to revert.",
    "restore_btn": "↩ Restore", "restore_ok": "Restored to previous state.",
    "trend_hdr": "📈 Onboarding Progress Over Time",
    "trend_no_gs": "Connect Google Sheets in the sidebar to view date-based trends.",
    "trend_spin": "Loading change history...",
    "trend_empty": "No history saved yet. Click '💾 Save current state' in the sidebar.",
    "trend_total": "**Change history** — total **{n} entries**",
    "footer": "VIS Issuer Onboarding Tracker · Powered by Claude AI · DealMe Internal",
}

def T(key, **kw):
    d = _EN if st.session_state.get("lang") == "en" else _KO
    s = d.get(key, _KO.get(key, key))
    return s.format(**kw) if kw else s

_STATUS_EN = {
    "완료": "Done", "미완료": "Not done", "진행중": "In progress",
    "확인중": "Checking", "확인필요": "Review needed", "심각지연": "Critical delay",
    "BIN 대기": "BIN pending", "BIN대기": "BIN pending", "미확인": "Unknown",
    "온보딩 시작 전": "Not started",
}
_TAG_EN = {"긴급": "Urgent", "확인필요": "Review", "대기중": "Pending", "완료": "Done"}

def ds(v):
    # Translate status value for display
    if st.session_state.get("lang") == "en":
        return _STATUS_EN.get(str(v), str(v))
    return str(v)

def dt(tag):
    # Translate action tag for display
    if st.session_state.get("lang") == "en":
        return _TAG_EN.get(tag, tag)
    return tag

# ── 페이지 설정 ────────────────────────────────────────────────────────────────
st.set_page_config(page_title="VIS 이슈어 온보딩 트래커", page_icon="💳", layout="wide")

st.markdown("""
<style>
  .tag-urgent  { background:#FCEBEB; color:#A32D2D; padding:2px 10px; border-radius:20px; font-size:12px; font-weight:600; }
  .tag-check   { background:#E6F1FB; color:#185FA5; padding:2px 10px; border-radius:20px; font-size:12px; font-weight:600; }
  .tag-pending { background:#FAEEDA; color:#854F0B; padding:2px 10px; border-radius:20px; font-size:12px; font-weight:600; }
  .tag-done    { background:#EAF3DE; color:#2D6A0E; padding:2px 10px; border-radius:20px; font-size:12px; font-weight:600; }
  .bin-chip    { display:inline-block; background:#F0F4FF; color:#1a3a8f; border:1px solid #C7D4F5;
                 padding:3px 10px; border-radius:16px; font-size:12px; font-family:monospace; margin:2px; }
</style>
""", unsafe_allow_html=True)


# ── 1. 비밀번호 로그인 ─────────────────────────────────────────────────────────
def check_password():
    try:
        correct = st.secrets["passwords"]["correct_password"]
    except Exception:
        correct = os.environ.get("APP_PASSWORD", "dealme2026")

    def _hash(pw): return hashlib.sha256(pw.encode()).hexdigest()

    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False
    if st.session_state.authenticated:
        return True

    st.markdown(f"""
    <div style="max-width:400px;margin:80px auto 0;text-align:center;">
      <div style="font-size:40px;margin-bottom:8px;">💳</div>
      <div style="font-size:22px;font-weight:700;margin-bottom:4px;">{T("login_title")}</div>
      <div style="font-size:13px;color:#888;margin-bottom:32px;">{T("login_sub")}</div>
    </div>
    """, unsafe_allow_html=True)

    _, col, _ = st.columns([1, 2, 1])
    with col:
        pw = st.text_input(T("login_pw"), type="password", placeholder=T("login_pw_ph"))
        if st.button(T("login_btn"), use_container_width=True, type="primary"):
            if _hash(pw) == _hash(correct):
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error(T("login_err"))
    st.stop()

check_password()


# ── 2. 상수 / 기본 데이터 ──────────────────────────────────────────────────────
# bin_list: 실제 BIN 번호 목록 (6~8자리 숫자 문자열 배열)
EMPTY_BANK = {
    "krw_vnd": "미완료", "usd": "미완료",
    "tc": "미완료",      "tc_url": "",
    "bin": "미완료",     "bin_list": [],
    "status": "미확인",  "note": "온보딩 시작 전",
}

DEFAULT_STATE = {
    "ACB": {
        "krw_vnd": "완료", "usd": "확인중", "tc": "완료",
        "tc_url": "https://acb.com.vn/acbwebsite/files/Ban_Dieu_khoan_dieu_kien_tra_gop_the_tin_dung_09.25.pdf",
        "bin": "완료", "bin_list": [],
        "status": "확인필요", "note": "USD 플랜 확인 중",
    },
    "Sacombank": {
        "krw_vnd": "완료", "usd": "완료", "tc": "완료",
        "tc_url": "https://www.sacombank.com.vn/content/dam/sacombank/files/the-le/sacombank-installment-en.pdf",
        "bin": "미완료", "bin_list": [],
        "status": "BIN 대기", "note": "BIN 리스트 미제출",
    },
    "TCB": {
        "krw_vnd": "완료", "usd": "진행중", "tc": "완료",
        "tc_url": "https://aemapp.techcombank.com/content/dam/techcombank/public-site/en/images/personal-banking/pl-06-hd-card-17-en-apendix-conditions-and-terms-of-use-of-credit-card-installments-a189549880-f3cae49033.pdf",
        "bin": "미완료", "bin_list": [],
        "status": "진행중", "note": "USD 세팅 중",
    },
    "Vietcombank": {
        "krw_vnd": "미완료", "usd": "미완료", "tc": "미완료", "tc_url": "",
        "bin": "미완료", "bin_list": [],
        "status": "심각지연", "note": "테스트 미통과",
    },
}

DEFAULT_ACTIONS = [
    {"tag": "확인필요", "text": "ACB — USD 플랜 최종 등록 완료 여부 Irene에게 확인 요청"},
    {"tag": "긴급",    "text": "Sacombank — Irene/Wisid에게 BIN 리스트 제출 촉구 필요"},
    {"tag": "대기중",  "text": "TCB — USD 플랜 세팅 완료 확인 후 BIN 리스트 수급 진행"},
    {"tag": "긴급",    "text": "Vietcombank — 로직 수정 이슈 단기 해결 불가. 신세계 런칭 시 제외 여부 의사결정 필요"},
]

_defaults = {
    "bank_state":   dict(DEFAULT_STATE),
    "bank_order":   list(DEFAULT_STATE.keys()),
    "actions":      list(DEFAULT_ACTIONS),
    "update_time":  "2026년 5월 15일 (이메일 기준)",
    "api_key":      "",
    "input_hashes": [],
    "history":      [],
}
for _k, _v in _defaults.items():
    if _k not in st.session_state:
        st.session_state[_k] = _v

# 기존 세션에 bin_list 필드가 없는 경우 자동 보완
for _bank in st.session_state.bank_state:
    if "bin_list" not in st.session_state.bank_state[_bank]:
        st.session_state.bank_state[_bank]["bin_list"] = []


# ── 3. 헬퍼 함수 ───────────────────────────────────────────────────────────────
def status_icon(s):
    return {"완료": "✅", "미완료": "❌", "진행중": "🔄", "확인중": "❓",
            "확인필요": "⚠️", "심각지연": "🚨", "BIN 대기": "⏳"}.get(s, "❓")

def card_bg(status):
    if "심각" in status or "지연" in status: return "#FFF0F0"
    if status == "완료":                      return "#F0FBF0"
    if status in ("미확인", ""):              return "#F8F9FA"
    if "확인필요" in status:                  return "#FFFBF0"
    return "#FFF8F0"

def card_border(status):
    if "심각" in status or "지연" in status: return "#FFAAAA"
    if status == "완료":                      return "#AADDAA"
    if "확인필요" in status:                  return "#FFD580"
    return "#E5E7EB"

def make_hash(text: str) -> str:
    return hashlib.sha256(text.strip().encode()).hexdigest()[:16]

def is_duplicate(h: str) -> bool:
    return h in st.session_state.input_hashes

def save_snapshot(label: str):
    snap = {
        "label":       label,
        "time":        datetime.now().strftime("%m/%d %H:%M"),
        "bank_state":  copy.deepcopy(st.session_state.bank_state),
        "bank_order":  list(st.session_state.bank_order),
        "actions":     copy.deepcopy(st.session_state.actions),
        "update_time": st.session_state.update_time,
    }
    st.session_state.history.append(snap)
    if len(st.session_state.history) > 10:
        st.session_state.history.pop(0)

def restore_snapshot(idx: int):
    snap = st.session_state.history[idx]
    st.session_state.bank_state  = copy.deepcopy(snap["bank_state"])
    st.session_state.bank_order  = list(snap["bank_order"])
    st.session_state.actions     = copy.deepcopy(snap["actions"])
    st.session_state.update_time = snap["update_time"]

def merge_bin_list(existing: list, new_bins: list) -> list:
    """기존 BIN 목록에 새 BIN을 중복 없이 추가."""
    merged = list(existing)
    for b in new_bins:
        b_str = str(b).strip()
        if b_str and b_str not in merged:
            merged.append(b_str)
    return sorted(merged)


# ── 4. AI 파싱 엔진 ────────────────────────────────────────────────────────────
_BUSINESS_RULES = """
[핵심 비즈니스 규칙 — 반드시 준수]
1. 최신 정보 우선: 이메일 스레드는 아래로 갈수록 오래된 인용입니다.
   반드시 '가장 최근(상단) 메시지'의 결론을 최종값으로 채택하고, 하단 인용문은 참고만 하세요.
2. 롤백 인지: A→B→A처럼 원복되는 흐름이 있으면 최종값은 A입니다. 중간값 B를 최종값으로 착각하지 마세요.
3. 특이사항 자동 감지:
   - 8자리 BIN 번호(예: 43116707)가 언급되면 → status='확인필요', 액션 아이템에 추가
   - T&C 업로드 오류 / 시스템 에러 문맥 → status='확인필요', 액션 아이템에 추가
   - 테스트 실패, 로직 오류 → status='심각지연'
4. 은행 약자 매핑: VCB=Vietcombank, SCB=Sacombank, TCB=TCB(Techcombank)
5. 필드 허용값:
   - krw_vnd / usd / bin: 완료 | 미완료 | 진행중 | 확인중
   - tc: 완료 | 미완료
   - status: 완료 | 진행중 | 확인필요 | BIN대기 | 심각지연 | 미확인
   - bin_list: 6~8자리 숫자 BIN 번호 문자열 배열. 언급된 모든 BIN을 추출하세요.
6. 온보딩과 무관한 내용이면 반드시 {"irrelevant": true, "reason": "..."} 형태로만 응답하세요.
7. 언급이 없는 은행 필드는 무조건 빈 문자열("") 또는 빈 배열([])로 남기세요.
"""

def _bank_json_template() -> str:
    lines = [
        f'  "{b}": {{"krw_vnd":"","usd":"","tc":"","tc_url":"","bin":"","bin_list":[],"status":"","note":""}}'
        for b in st.session_state.bank_order
    ]
    return "{\n" + ",\n".join(lines) + "\n}"

def _call_claude_text(prompt: str) -> dict:
    client = anthropic.Anthropic(api_key=st.session_state.api_key)
    msg = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4000,
        messages=[{"role": "user", "content": prompt}],
    )
    raw = re.sub(r"```json|```", "", msg.content[0].text).strip()
    return json.loads(raw)

def _call_claude_vision(image_bytes: bytes, media_type: str, prompt: str) -> dict:
    client = anthropic.Anthropic(api_key=st.session_state.api_key)
    msg = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4000,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {
                    "type": "base64", "media_type": media_type,
                    "data": base64.b64encode(image_bytes).decode(),
                }},
                {"type": "text", "text": prompt},
            ],
        }],
    )
    raw = re.sub(r"```json|```", "", msg.content[0].text).strip()
    return json.loads(raw)

def analyze_email(text: str) -> dict:
    prompt = f"""{_BUSINESS_RULES}

[작업]
아래 이메일 스레드를 읽고 VIS 이슈어 온보딩 현황을 분석하세요.
등록 은행: {', '.join(st.session_state.bank_order)}

- tc_url: 이메일에 T&C 원문 PDF/웹 URL이 있으면 해당 은행의 tc_url에 넣으세요.
- bin_list: 본문에 등장하는 모든 6~8자리 BIN 번호를 배열로 추출하세요. 없으면 [].
- actions: 후속 조치 항목 최대 5개. tag는 긴급|확인필요|대기중 중 하나.
- update_date: 확인되는 가장 최근 날짜 "YYYY년 MM월 DD일" 형식.

[이메일 본문]
{text}

[출력 — JSON만 반환, 다른 텍스트 절대 금지]
관련 없음: {{"irrelevant": true, "reason": "이유"}}
관련 있음:
{{
  "irrelevant": false,
  "banks": {_bank_json_template()},
  "new_banks": [],
  "actions": [{{"tag": "긴급|확인필요|대기중", "text": ""}}],
  "update_date": ""
}}"""
    return _call_claude_text(prompt)

def analyze_excel(sheet_data: dict) -> dict:
    content = ""
    for sname, rows in sheet_data.items():
        content += f"\n[시트: {sname}]\n"
        content += "".join("\t".join(r) + "\n" for r in rows[:80])

    prompt = f"""{_BUSINESS_RULES}

[작업]
아래는 VIS 온보딩 관련 엑셀 파일 내용입니다.
등록 은행: {', '.join(st.session_state.bank_order)}

- tc_url: 엑셀에 T&C URL이 있으면 해당 은행의 tc_url에 넣으세요.
- bin_list: 파일에 등장하는 모든 6~8자리 BIN 번호를 배열로 추출하세요. 없으면 [].
- bin_count: 파악된 BIN 번호 총 개수.
- submitting_bank: 이 파일을 제출한 은행 이름.

[엑셀 내용]
{content}

[출력 — JSON만 반환]
관련 없음: {{"irrelevant": true, "reason": "이유"}}
관련 있음:
{{
  "irrelevant": false,
  "banks": {_bank_json_template()},
  "new_banks": [],
  "actions": [{{"tag": "긴급|확인필요|대기중", "text": ""}}],
  "update_date": "",
  "submitting_bank": "",
  "bin_count": 0
}}"""
    return _call_claude_text(prompt)

def analyze_image(image_bytes: bytes, media_type: str) -> dict:
    prompt = f"""{_BUSINESS_RULES}

[작업]
이 이미지는 VIS 이슈어 온보딩 관련 문서나 이메일 캡처일 수 있습니다.
등록 은행: {', '.join(st.session_state.bank_order)}

- tc_url: 이미지에 T&C URL이 보이면 반드시 추출하세요.
- bin_list: 이미지에서 보이는 모든 6~8자리 BIN 번호를 배열로 추출하세요. 없으면 [].

[출력 — JSON만 반환]
관련 없는 이미지: {{"irrelevant": true, "reason": "관련없는 이미지입니다"}}
관련 있는 이미지:
{{
  "irrelevant": false,
  "banks": {_bank_json_template()},
  "new_banks": [],
  "actions": [{{"tag": "긴급|확인필요|대기중", "text": ""}}],
  "update_date": ""
}}"""
    return _call_claude_vision(image_bytes, media_type, prompt)


# ── 5. 상태 업데이트 로직 ──────────────────────────────────────────────────────
def read_excel_bytes(file_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    out = {}
    for name in wb.sheetnames:
        rows = []
        for row in wb[name].iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append([str(c) if c is not None else "" for c in row])
        out[name] = rows
    return out

def apply_update(parsed: dict, input_hash: str = None):
    """파싱 결과를 session_state에 반영. irrelevant면 완전 무시."""
    if parsed.get("irrelevant"):
        return

    save_snapshot("업데이트 전")
    if input_hash:
        st.session_state.input_hashes.append(input_hash)

    # 새 은행 추가
    for nb in parsed.get("new_banks", []):
        if nb and nb not in st.session_state.bank_state:
            st.session_state.bank_state[nb] = dict(EMPTY_BANK)
            st.session_state.bank_state[nb]["bin_list"] = []
            st.session_state.bank_order.append(nb)

    # 은행 상태 업데이트
    for bank, info in parsed.get("banks", {}).items():
        if not info:
            continue
        if bank not in st.session_state.bank_state:
            st.session_state.bank_state[bank] = dict(EMPTY_BANK)
            st.session_state.bank_state[bank]["bin_list"] = []
            st.session_state.bank_order.append(bank)

        for k, v in info.items():
            if k == "bin_list":
                # BIN 목록은 누적 병합 (덮어쓰지 않음)
                if isinstance(v, list) and v:
                    existing = st.session_state.bank_state[bank].get("bin_list", [])
                    st.session_state.bank_state[bank]["bin_list"] = merge_bin_list(existing, v)
            elif v:
                st.session_state.bank_state[bank][k] = v

    # 액션 아이템: 실질 내용 있을 때만 교체
    new_actions = [a for a in parsed.get("actions", []) if a.get("text", "").strip()]
    if new_actions:
        st.session_state.actions = new_actions

    if parsed.get("update_date"):
        st.session_state.update_time = parsed["update_date"]


# ── 6. Google Sheets 연동 ──────────────────────────────────────────────────────
_SCOPES = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
_FIELD_LABELS = {
    "krw_vnd": "KRW/VND 플랜", "usd": "USD 플랜",
    "tc": "영어 T&C", "bin": "BIN 리스트",
    "status": "전체 상태", "note": "비고",
}

def get_gsheet_client():
    try:
        creds = Credentials.from_service_account_info(
            dict(st.secrets["gcp_service_account"]), scopes=_SCOPES
        )
        return gspread.authorize(creds)
    except Exception:
        return None

def save_to_gsheet(client, spreadsheet_id: str):
    try:
        wb    = client.open_by_key(spreadsheet_id)
        banks = st.session_state.bank_order
        state = st.session_state.bank_state
        now   = datetime.now().strftime("%Y-%m-%d %H:%M")

        # dashboard 시트
        try:
            ws_dash = wb.worksheet("dashboard")
            ws_dash.clear()
        except gspread.WorksheetNotFound:
            ws_dash = wb.add_worksheet(title="dashboard", rows=60, cols=12)

        header = ["은행", "KRW/VND", "USD", "T&C", "T&C URL", "BIN상태", "BIN수", "전체상태", "비고", "마지막업데이트"]
        rows = [header]
        for bank in banks:
            info = state.get(bank, dict(EMPTY_BANK))
            bin_cnt = len(info.get("bin_list", []))
            rows.append([
                bank,
                info.get("krw_vnd", ""), info.get("usd", ""),
                info.get("tc", ""),      info.get("tc_url", ""),
                info.get("bin", ""),     bin_cnt,
                info.get("status", ""),  info.get("note", ""), now,
            ])
        ws_dash.update("A1", rows)
        ws_dash.format("A1:J1", {
            "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
            "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
        })

        # bin_list 시트
        try:
            ws_bin = wb.worksheet("bin_list")
            ws_bin.clear()
        except gspread.WorksheetNotFound:
            ws_bin = wb.add_worksheet(title="bin_list", rows=2000, cols=3)
        bin_header = ["은행", "BIN 번호", "저장일시"]
        bin_rows = [bin_header]
        for bank in banks:
            info = state.get(bank, dict(EMPTY_BANK))
            for b in info.get("bin_list", []):
                bin_rows.append([bank, b, now])
        ws_bin.update("A1", bin_rows)
        ws_bin.format("A1:C1", {
            "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
            "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
        })

        # history 시트
        try:
            ws_hist = wb.worksheet("history")
        except gspread.WorksheetNotFound:
            ws_hist = wb.add_worksheet(title="history", rows=2000, cols=6)
            ws_hist.append_row(["날짜", "은행", "항목", "이전값", "현재값", "비고"])
            ws_hist.format("A1:F1", {
                "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                "backgroundColor": {"red": 0.12, "green": 0.31, "blue": 0.47},
            })

        last_state: dict = {}
        for rec in ws_hist.get_all_records():
            b, f, v = rec.get("은행", ""), rec.get("항목", ""), rec.get("현재값", "")
            if b and f:
                last_state.setdefault(b, {})[f] = v

        new_rows = []
        for bank in banks:
            info = state.get(bank, dict(EMPTY_BANK))
            for field in ["krw_vnd", "usd", "tc", "bin", "status"]:
                label    = _FIELD_LABELS[field]
                cur_val  = info.get(field, "")
                prev_val = last_state.get(bank, {}).get(label)
                if prev_val is None or prev_val != cur_val:
                    new_rows.append([now, bank, label,
                                     prev_val if prev_val else "신규",
                                     cur_val, info.get("note", "")])
        if new_rows:
            ws_hist.append_rows(new_rows)

        return True, len(new_rows)
    except Exception as e:
        return False, str(e)

def load_history_from_gsheet(client, spreadsheet_id: str) -> list:
    try:
        return client.open_by_key(spreadsheet_id).worksheet("history").get_all_records()
    except Exception:
        return []


# ── 7. 엑셀 내보내기 ───────────────────────────────────────────────────────────
def export_excel_file() -> BytesIO:
    wb = openpyxl.Workbook()

    hdr_fill  = PatternFill("solid", fgColor="1F4E79")
    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    center    = Alignment(horizontal="center", vertical="center")
    wrap_left = Alignment(vertical="center", wrap_text=True)
    thin      = Side(style="thin", color="CCCCCC")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)
    sc = {"완료": "EAF3DE", "미완료": "FCEBEB", "진행중": "FFF3CD",
          "확인중": "E6F1FB", "확인필요": "FFF3CD", "해당없음": "F1EFE8"}

    def write_header(ws, headers, col_widths):
        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, border
            ws.column_dimensions[c.column_letter].width = w

    # 시트1: 온보딩 현황
    ws1 = wb.active
    ws1.title = "온보딩 현황"
    write_header(ws1,
                 ["은행명", "KRW/VND 플랜", "USD 플랜", "영어 T&C", "BIN 리스트", "BIN 수", "전체 상태", "비고"],
                 [15, 16, 14, 14, 14, 10, 14, 35])
    for r, bank in enumerate(st.session_state.bank_order, 2):
        info    = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        bin_cnt = len(info.get("bin_list", []))
        vals = [bank, info["krw_vnd"], info["usd"], info["tc"],
                info["bin"], bin_cnt, info["status"], info["note"]]
        for c_idx, val in enumerate(vals, 1):
            cell = ws1.cell(row=r, column=c_idx, value=val)
            cell.border = border
            cell.alignment = center if c_idx != 8 else wrap_left
            if c_idx in (2, 3, 4, 5, 7) and str(val) in sc:
                cell.fill = PatternFill("solid", fgColor=sc[str(val)])
        ws1.row_dimensions[r].height = 22

    # 시트2: BIN 리스트 (은행별)
    ws_bin = wb.create_sheet("BIN 리스트")
    write_header(ws_bin, ["은행명", "BIN 번호", "자리수"], [15, 14, 10])
    bin_row = 2
    for bank in st.session_state.bank_order:
        info     = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        bin_list = info.get("bin_list", [])
        if bin_list:
            for b in bin_list:
                ws_bin.cell(row=bin_row, column=1, value=bank).border = border
                bc = ws_bin.cell(row=bin_row, column=2, value=b)
                bc.border, bc.alignment = border, center
                bc.font = Font(name="Courier New", size=11)
                dc = ws_bin.cell(row=bin_row, column=3, value=len(str(b)))
                dc.border, dc.alignment = border, center
                ws_bin.row_dimensions[bin_row].height = 20
                bin_row += 1
        else:
            ws_bin.cell(row=bin_row, column=1, value=bank).border = border
            nc = ws_bin.cell(row=bin_row, column=2, value="(미제출)")
            nc.border, nc.alignment, nc.font = border, center, Font(color="999999", italic=True)
            ws_bin.cell(row=bin_row, column=3, value="").border = border
            ws_bin.row_dimensions[bin_row].height = 20
            bin_row += 1

    # BIN 은행별 요약 (우측)
    ws_bin.cell(row=1, column=5, value="은행별 BIN 수 요약").font = Font(bold=True)
    for sr, bank in enumerate(st.session_state.bank_order, 2):
        info = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        ws_bin.cell(row=sr, column=5, value=bank)
        ws_bin.cell(row=sr, column=6, value=len(info.get("bin_list", []))).alignment = center
    ws_bin.column_dimensions["E"].width = 15
    ws_bin.column_dimensions["F"].width = 10

    # 시트3: T&C 현황
    ws_tc = wb.create_sheet("T&C 현황")
    write_header(ws_tc, ["은행명", "T&C 상태", "원문 URL"], [15, 12, 90])
    for r, bank in enumerate(st.session_state.bank_order, 2):
        info  = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        tc_v  = info.get("tc", "미완료")
        tc_u  = info.get("tc_url", "")
        ws_tc.cell(row=r, column=1, value=bank).border = border
        tc_c = ws_tc.cell(row=r, column=2, value=tc_v)
        tc_c.alignment, tc_c.border = center, border
        if tc_v in sc:
            tc_c.fill = PatternFill("solid", fgColor=sc[tc_v])
        url_c = ws_tc.cell(row=r, column=3, value=tc_u)
        url_c.alignment, url_c.border = wrap_left, border
        if tc_u:
            url_c.font = Font(color="185FA5", underline="single")
        ws_tc.row_dimensions[r].height = 22

    # 시트4: 액션 아이템
    ws_act = wb.create_sheet("액션 아이템")
    write_header(ws_act, ["태그", "내용"], [14, 80])
    tag_c = {"긴급": "FCEBEB", "확인필요": "E6F1FB", "대기중": "FAEEDA"}
    for r, a in enumerate(st.session_state.actions, 2):
        t = ws_act.cell(row=r, column=1, value=a["tag"])
        t.alignment, t.border = center, border
        t.fill = PatternFill("solid", fgColor=tag_c.get(a["tag"], "FFFFFF"))
        n = ws_act.cell(row=r, column=2, value=a["text"])
        n.alignment, n.border = wrap_left, border
        ws_act.row_dimensions[r].height = 22

    # 시트5: 요약
    ws_sum = wb.create_sheet("요약")
    ws_sum.column_dimensions["A"].width = 22
    ws_sum.column_dimensions["B"].width = 18
    bstate  = st.session_state.bank_state
    total   = len(bstate)
    ready   = sum(1 for b in bstate.values() if b["status"] == "완료")
    crit    = sum(1 for b in bstate.values() if "심각" in b["status"])
    tc_done = sum(1 for b in bstate.values() if b.get("tc") == "완료")
    total_bins = sum(len(b.get("bin_list", [])) for b in bstate.values())
    for r, (k, v) in enumerate([
        ("내보내기 일시",    datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("마지막 업데이트", st.session_state.update_time),
        ("전체 은행 수",     total),
        ("런칭 준비 완료",  ready),
        ("진행 중",          total - ready - crit),
        ("심각한 지연",     crit),
        ("T&C 완료",         tc_done),
        ("등록된 BIN 총계",  total_bins),
    ], 1):
        ws_sum.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws_sum.cell(row=r, column=2, value=v)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── 8. PDF 내보내기 ────────────────────────────────────────────────────────────
def export_pdf_file() -> BytesIO:
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    KO = "HYGothic-Medium"

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)

    s_title = ParagraphStyle("t", fontName=KO, fontSize=16, spaceAfter=4, leading=22)
    s_sub   = ParagraphStyle("s", fontName=KO, fontSize=9,  textColor=colors.grey, spaceAfter=12, leading=14)
    s_h2    = ParagraphStyle("h", fontName=KO, fontSize=12, spaceBefore=14, spaceAfter=6, leading=18)

    sbg = {"완료": colors.HexColor("#EAF3DE"), "미완료": colors.HexColor("#FCEBEB"),
           "진행중": colors.HexColor("#FFF3CD"), "확인중": colors.HexColor("#E6F1FB"),
           "확인필요": colors.HexColor("#FFF3CD")}
    tbg = {"긴급": colors.HexColor("#FCEBEB"), "확인필요": colors.HexColor("#E6F1FB"),
           "대기중": colors.HexColor("#FAEEDA")}

    def base_ts(extra=None):
        s = [
            ("FONTNAME",   (0, 0), (-1, -1), KO),
            ("BACKGROUND", (0, 0), (-1, 0),  colors.HexColor("#1F4E79")),
            ("TEXTCOLOR",  (0, 0), (-1, 0),  colors.white),
            ("FONTSIZE",   (0, 0), (-1, -1), 8),
            ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
            ("VALIGN",     (0, 0), (-1, -1), "MIDDLE"),
            ("GRID",       (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
            ("ROWHEIGHT",  (0, 0), (-1, -1), 16),
        ]
        if extra:
            s.extend(extra)
        return TableStyle(s)

    story = []
    story.append(Paragraph("VIS 이슈어 온보딩 현황 보고서", s_title))
    story.append(Paragraph(
        f"신세계면세점 런칭 기준 | 업데이트: {st.session_state.update_time} "
        f"| 출력: {datetime.now().strftime('%Y-%m-%d %H:%M')}", s_sub))

    bstate     = st.session_state.bank_state
    total      = len(bstate)
    ready      = sum(1 for b in bstate.values() if b["status"] == "완료")
    crit       = sum(1 for b in bstate.values() if "심각" in b["status"])
    total_bins = sum(len(b.get("bin_list", [])) for b in bstate.values())

    kpi = Table(
        [["전체 은행", "준비 완료", "진행 중", "등록 BIN"],
         [str(total), str(ready), str(total - ready - crit), str(total_bins)]],
        colWidths=[44*mm]*4)
    kpi.setStyle(base_ts([("FONTSIZE", (0,1),(-1,1),14),("ROWHEIGHT",(0,1),(-1,1),24)]))
    story.append(kpi)
    story.append(Spacer(1, 8*mm))

    # 은행별 현황
    story.append(Paragraph("은행별 온보딩 현황", s_h2))
    fl = [("krw_vnd","KRW/VND"),("usd","USD"),("tc","T&C"),("bin","BIN")]
    tdata = [["은행명","KRW/VND","USD","T&C","BIN","BIN수","상태","비고"]]
    ext = []
    for ri, bank in enumerate(st.session_state.bank_order, 1):
        info    = bstate.get(bank, dict(EMPTY_BANK))
        bin_cnt = len(info.get("bin_list", []))
        tdata.append([bank]+[info[f] for f,_ in fl]+[str(bin_cnt), info["status"], info["note"]])
        for ci, (fk,_) in enumerate(fl, 1):
            if info[fk] in sbg:
                ext.append(("BACKGROUND",(ci,ri),(ci,ri),sbg[info[fk]]))
        if info["status"] in sbg:
            ext.append(("BACKGROUND",(6,ri),(6,ri),sbg[info["status"]]))
        ext.append(("ALIGN",(7,ri),(7,ri),"LEFT"))
    t2 = Table(tdata, colWidths=[20*mm,16*mm,14*mm,12*mm,18*mm,10*mm,18*mm,54*mm])
    t2.setStyle(base_ts(ext))
    story.append(t2)
    story.append(Spacer(1, 8*mm))

    # BIN 리스트
    story.append(Paragraph("은행별 BIN 리스트", s_h2))
    bdata = [["은행명", "BIN 번호 목록", "수량"]]
    for bank in st.session_state.bank_order:
        info     = bstate.get(bank, dict(EMPTY_BANK))
        bin_list = info.get("bin_list", [])
        bin_str  = "  ".join(bin_list) if bin_list else "(미제출)"
        bdata.append([bank, bin_str, str(len(bin_list))])
    t3 = Table(bdata, colWidths=[25*mm, 130*mm, 12*mm])
    t3.setStyle(base_ts([("ALIGN",(1,1),(1,-1),"LEFT")]))
    story.append(t3)
    story.append(Spacer(1, 8*mm))

    # 액션 아이템
    story.append(Paragraph("액션 아이템", s_h2))
    adata = [["태그","내용"]]
    aext  = []
    for ri, a in enumerate(st.session_state.actions, 1):
        adata.append([a["tag"], a["text"]])
        if a["tag"] in tbg:
            aext.append(("BACKGROUND",(0,ri),(0,ri),tbg[a["tag"]]))
        aext.extend([("ROWHEIGHT",(0,ri),(-1,ri),22),("ALIGN",(1,ri),(1,ri),"LEFT")])
    t4 = Table(adata, colWidths=[24*mm, 152*mm])
    t4.setStyle(base_ts(aext))
    story.append(t4)

    doc.build(story)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# UI
# ══════════════════════════════════════════════════════════════════════════════

# ── 사이드바 ──────────────────────────────────────────────────────────────────
with st.sidebar:
    if st.button(T("lang_toggle"), key="lang_btn", use_container_width=True):
        st.session_state["lang"] = "en" if st.session_state.get("lang") == "ko" else "ko"
        st.rerun()
    st.divider()
    st.header(T("settings"))
    key_in = st.text_input("Anthropic API Key", type="password",
                           value=st.session_state.api_key, placeholder="sk-ant-...")
    if key_in:
        st.session_state.api_key = key_in
    if st.session_state.api_key:
        st.success(T("api_key_set"))
    else:
        st.warning(T("api_key_warn"))

    st.divider()
    st.header(T("gsheet_hdr"))
    try:
        _sheet_id = st.secrets["google"]["spreadsheet_id"]
        st.caption(f"시트 ID: `{_sheet_id[:18]}...`")
    except Exception:
        _sheet_id = st.text_input("Spreadsheet ID", placeholder=T("sheet_id_ph"))
    st.session_state["_sheet_id"] = _sheet_id if _sheet_id else ""

    _gs_client = get_gsheet_client()
    if _gs_client and _sheet_id:
        st.success(T("gsheet_ok"))
        if st.button(T("save_btn")):
            with st.spinner("저장 중..."):
                ok, result = save_to_gsheet(_gs_client, _sheet_id)
                if ok:
                    st.success(T("save_ok", n=result))
                else:
                    st.error(T("save_fail", e=result))
    else:
        st.warning(T("gsheet_warn"))

    st.divider()
    st.header(T("bank_mgmt"))
    new_bank = st.text_input(T("add_bank_label"), placeholder=T("add_bank_ph"))
    if st.button(T("add_btn")):
        n = new_bank.strip()
        if not n:
            st.warning(T("add_warn_empty"))
        elif n in st.session_state.bank_state:
            st.warning(T("add_warn_dup", n=n))
        else:
            st.session_state.bank_state[n] = dict(EMPTY_BANK)
            st.session_state.bank_state[n]["bin_list"] = []
            st.session_state.bank_order.append(n)
            st.success(T("add_ok", n=n))
            st.rerun()

    del_target = st.selectbox(T("del_bank_label"), [T("del_select")] + st.session_state.bank_order)
    if st.button(T("del_btn")):
        if del_target != T("del_select"):
            st.session_state.bank_state.pop(del_target, None)
            if del_target in st.session_state.bank_order:
                st.session_state.bank_order.remove(del_target)
            st.success(T("del_ok", n=del_target))
            st.rerun()

    st.divider()
    st.markdown(T("legend"))
    st.divider()
    c_lo, c_re = st.columns(2)
    with c_lo:
        if st.button(T("logout_btn"), use_container_width=True):
            st.session_state.authenticated = False
            st.rerun()
    with c_re:
        if st.button(T("reset_btn"), use_container_width=True):
            for k in ["bank_state","bank_order","actions","update_time","input_hashes","history"]:
                st.session_state.pop(k, None)
            st.rerun()


# ── 메인 헤더 ──────────────────────────────────────────────────────────────────
st.title(T("app_title"))
st.caption(T("app_sub", t=st.session_state.update_time))

# KPI
_st      = st.session_state.bank_state
_total   = len(_st)
_ready   = sum(1 for b in _st.values() if b["status"] == "완료")
_crit    = sum(1 for b in _st.values() if "심각" in b["status"])
_tc_done = sum(1 for b in _st.values() if b.get("tc") == "완료")
_tc_url  = sum(1 for b in _st.values() if b.get("tc_url"))
_total_bins = sum(len(b.get("bin_list", [])) for b in _st.values())

k1,k2,k3,k4,k5,k6,k7 = st.columns(7)
k1.metric(T("kpi_total"), _total)
k2.metric(T("kpi_ready"), _ready)
k3.metric(T("kpi_inprog"), _total - _ready - _crit)
k4.metric(T("kpi_crit"), _crit)
k5.metric(T("kpi_tc"), _tc_done)
k6.metric(T("kpi_url"), _tc_url)
k7.metric(T("kpi_bin"), _total_bins)

st.divider()

# 내보내기
with st.expander(T("export_exp"), expanded=False):
    ex1, ex2 = st.columns(2)
    with ex1:
        buf_x = export_excel_file()
        st.download_button(T("export_xlsx"), data=buf_x,
                           file_name=f"VIS_온보딩_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           use_container_width=True)
    with ex2:
        buf_p = export_pdf_file()
        st.download_button(T("export_pdf"), data=buf_p,
                           file_name=f"VIS_온보딩_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                           mime="application/pdf", use_container_width=True)


# ── 대시보드 탭 (온보딩 현황 / T&C 현황 / BIN 리스트) ─────────────────────────
dash_tab1, dash_tab2, dash_tab3 = st.tabs([T("tab_status"), T("tab_tc"), T("tab_bin")])

# ── 탭1: 온보딩 현황 ───────────────────────────────────────────────────────────
with dash_tab1:
    st.subheader(T("tab1_hdr"))
    _fields = [("krw_vnd", T("fld_krw")), ("usd", T("fld_usd")), ("tc", T("fld_tc")), ("bin", T("fld_bin"))]
    cols_card = st.columns(2)

    for i, bank in enumerate(st.session_state.bank_order):
        if bank not in _st:
            continue
        info   = _st[bank]
        tc_url = info.get("tc_url", "")
        tc_lnk = (f' <a href="{tc_url}" target="_blank" style="font-size:11px;color:#185FA5;">📄 원문</a>') if tc_url else ""

        rows_html = ""
        for fk, lbl in _fields:
            extra = tc_lnk if fk == "tc" else ""
            _bl = info.get("bin_list", [])
            bin_cnt_str = f" <span style='font-size:11px;color:#888;'>({len(_bl)}개)</span>" if fk == "bin" else ""
            rows_html += (
                f'<div style="display:flex;justify-content:space-between;'
                f'padding:5px 0;font-size:13px;border-bottom:1px solid rgba(0,0,0,0.06);">'
                f'<span style="color:#555;">{lbl}</span>'
                f'<span style="font-weight:500;">{status_icon(info[fk])} {info[fk]}{bin_cnt_str}{extra}</span></div>'
            )

        with cols_card[i % 2]:
            st.markdown(f"""
            <div style="background:{card_bg(info['status'])};border-radius:14px;padding:18px 22px;
                        margin-bottom:14px;border:1.5px solid {card_border(info['status'])};">
              <div style="font-size:17px;font-weight:700;margin-bottom:10px;">
                {bank}
                <span style="font-size:12px;font-weight:400;color:#777;margin-left:8px;">
                  {status_icon(info['status'])} {info['status']}
                </span>
              </div>
              {rows_html}
              <div style="font-size:12px;color:#888;margin-top:8px;">📝 {info.get('note','')}</div>
            </div>""", unsafe_allow_html=True)

            with st.expander(T("edit_exp", b=bank)):
                _opts = {
                    "krw_vnd": ["완료","미완료","진행중","확인중"],
                    "usd":     ["완료","미완료","진행중","확인중"],
                    "tc":      ["완료","미완료"],
                    "bin":     ["완료","미완료","진행중","확인중"],
                    "status":  ["완료","진행중","확인필요","BIN 대기","심각지연","미확인"],
                }
                edited = {}
                for fk, opts in _opts.items():
                    cur = info.get(fk, opts[0])
                    idx = opts.index(cur) if cur in opts else 0
                    edited[fk] = st.selectbox(fk.upper().replace("_","/"), opts,
                                              index=idx, key=f"edit_{bank}_{fk}")
                edited["note"]   = st.text_input("비고", value=info.get("note",""), key=f"edit_{bank}_note")
                edited["tc_url"] = st.text_input("T&C URL", value=info.get("tc_url",""),
                                                 key=f"edit_{bank}_tc_url", placeholder="https://...")
                if st.button(T("save_edit_btn"), key=f"save_{bank}"):
                    save_snapshot(f"{bank} 수동편집 전")
                    for fk, val in edited.items():
                        st.session_state.bank_state[bank][fk] = val
                    st.success(T("save_edit_ok", b=bank))
                    st.rerun()

    st.divider()
    st.subheader(T("actions_hdr"))
    _tag_cls = {"긴급":"tag-urgent","확인필요":"tag-check","대기중":"tag-pending","완료":"tag-done","Urgent":"tag-urgent","Review":"tag-check","Pending":"tag-pending","Done":"tag-done"}
    if st.session_state.actions:
        for a in st.session_state.actions:
            cls = _tag_cls.get(a.get("tag",""), "tag-check")
            st.markdown(f'<span class="{cls}">{dt(a["tag"])}</span>&nbsp;&nbsp;{a["text"]}',
                        unsafe_allow_html=True)
    else:
        st.info(T("actions_empty"))


# ── 탭2: T&C 현황 ─────────────────────────────────────────────────────────────
with dash_tab2:
    st.subheader(T("tab2_hdr"))
    st.caption(T("tab2_cap"))

    m1, m2, m3 = st.columns(3)
    m1.metric(T("tc_m_total"), len(st.session_state.bank_order))
    m2.metric(T("tc_m_done"), _tc_done)
    m3.metric(T("tc_m_url"), _tc_url)

    st.markdown("---")
    _h1,_h2,_h3,_h4 = st.columns([2,1.2,1,5])
    _h1.markdown(T("tc_col_bank")); _h2.markdown(T("tc_col_stat"))
    _h3.markdown(T("tc_col_url")); _h4.markdown(T("tc_col_link"))
    st.markdown("---")

    for bank in st.session_state.bank_order:
        info      = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        tc_status = info.get("tc","미완료")
        tc_url    = info.get("tc_url","")
        c1,c2,c3,c4 = st.columns([2,1.2,1,5])
        c1.markdown(f"**{bank}**")
        c2.markdown(f"{status_icon(tc_status)} {tc_status}")
        c3.markdown("✅" if tc_url else "❌")
        if tc_url:
            c4.markdown(f"[{T('tc_view')}]({tc_url})")
        else:
            c4.caption(T("tc_no_url"))

        with st.expander(T("tc_edit_exp", b=bank)):
            new_url = st.text_input(T("tc_url_label"), value=tc_url,
                                    key=f"tc_url_{bank}", placeholder="https://...")
            if st.button(T("tc_save_btn"), key=f"tc_save_{bank}"):
                save_snapshot(f"{bank} T&C 수정 전")
                st.session_state.bank_state[bank]["tc_url"] = new_url.strip()
                if new_url.strip():
                    st.session_state.bank_state[bank]["tc"] = "완료"
                st.success(T("tc_save_ok", b=bank))
                st.rerun()


# ── 탭3: BIN 리스트 ───────────────────────────────────────────────────────────
with dash_tab3:
    st.subheader(T("tab3_hdr"))
    st.caption(T("tab3_cap"))

    # 전체 요약 메트릭
    b1, b2, b3 = st.columns(3)
    _banks_with_bin = sum(1 for b in _st.values() if b.get("bin_list"))
    b1.metric(T("bin_m_with"), _banks_with_bin)
    b2.metric(T("bin_m_without"), len(_st) - _banks_with_bin)
    b3.metric(T("bin_m_total"), _total_bins)

    st.markdown("---")

    for bank in st.session_state.bank_order:
        info     = st.session_state.bank_state.get(bank, dict(EMPTY_BANK))
        bin_list = info.get("bin_list", [])
        bin_cnt  = len(bin_list)

        # 은행 헤더
        col_title, col_badge = st.columns([5, 1])
        with col_title:
            st.markdown(f"### 🏦 {bank}")
        with col_badge:
            badge_color = "#EAF3DE" if bin_cnt > 0 else "#FCEBEB"
            badge_text  = "#2D6A0E" if bin_cnt > 0 else "#A32D2D"
            st.markdown(
                f'<div style="background:{badge_color};color:{badge_text};'
                f'border-radius:20px;padding:4px 14px;text-align:center;'
                f'font-weight:700;font-size:14px;margin-top:8px;">'
                f'{bin_cnt}개</div>',
                unsafe_allow_html=True,
            )

        if bin_list:
            # BIN 칩 형태로 표시
            chips_html = "".join(f'<span class="bin-chip">{b}</span>' for b in bin_list)
            st.markdown(
                f'<div style="padding:10px 0 4px;">{chips_html}</div>',
                unsafe_allow_html=True,
            )

            # BIN 테이블 (expander)
            with st.expander(T("bin_exp", b=bank, n=bin_cnt)):
                # 테이블 형태로 표시
                col_headers = st.columns([1, 2, 2])
                col_headers[0].markdown(T("bin_col_num"))
                col_headers[1].markdown(T("bin_col_bin"))
                col_headers[2].markdown(T("bin_col_digits"))
                for idx, b in enumerate(bin_list, 1):
                    r1, r2, r3 = st.columns([1, 2, 2])
                    r1.write(idx)
                    r2.code(b)
                    r3.write(f"{len(str(b))}자리")

                st.markdown("---")
                # BIN 삭제
                del_bin = st.selectbox(T("bin_del_sel"), [T("del_select")] + bin_list,
                                       key=f"del_bin_{bank}")
                if st.button(T("bin_del_btn"), key=f"del_bin_btn_{bank}"):
                    if del_bin != T("del_select"):
                        save_snapshot(f"{bank} BIN 삭제 전")
                        st.session_state.bank_state[bank]["bin_list"].remove(del_bin)
                        st.success(T("bin_del_ok", b=del_bin))
                        st.rerun()

                # 전체 BIN 클리어
                if st.button(T("bin_clear_btn", b=bank), key=f"clear_bin_{bank}",
                             type="secondary"):
                    save_snapshot(f"{bank} BIN 전체삭제 전")
                    st.session_state.bank_state[bank]["bin_list"] = []
                    st.session_state.bank_state[bank]["bin"] = "미완료"
                    st.success(T("bin_clear_ok", b=bank))
                    st.rerun()
        else:
            st.caption(T("bin_none"))

        # BIN 수동 입력 (항상 표시)
        with st.expander(T("bin_manual_exp", b=bank)):
            st.caption(T("bin_manual_cap"))
            manual_bin_input = st.text_area(
                T("bin_manual_label"), height=100,
                key=f"manual_bin_{bank}",
                placeholder="431167\n43116707\n456789",
            )
            if st.button(T("bin_add_btn"), key=f"add_bin_{bank}", type="primary"):
                raw = manual_bin_input.replace(",", "\n")
                new_bins = [
                    b.strip() for b in raw.split("\n")
                    if b.strip().isdigit() and 6 <= len(b.strip()) <= 8
                ]
                invalid = [
                    b.strip() for b in raw.split("\n")
                    if b.strip() and not (b.strip().isdigit() and 6 <= len(b.strip()) <= 8)
                ]
                if new_bins:
                    save_snapshot(f"{bank} BIN 수동추가 전")
                    existing = st.session_state.bank_state[bank].get("bin_list", [])
                    merged   = merge_bin_list(existing, new_bins)
                    added    = len(merged) - len(existing)
                    st.session_state.bank_state[bank]["bin_list"] = merged
                    if merged:
                        st.session_state.bank_state[bank]["bin"] = "완료"
                    st.success(T("bin_add_ok", added=added, total=len(merged)))
                if invalid:
                    st.warning(T("bin_invalid", items=', '.join(invalid)))
                if new_bins:
                    st.rerun()

        st.markdown("---")


# ── 입력 섹션 ─────────────────────────────────────────────────────────────────
st.subheader(T("input_hdr"))
st.caption(T("input_cap"))

inp1, inp2, inp3 = st.tabs([T("inp_email"), T("inp_excel"), T("inp_img")])

with inp1:
    st.caption(T("email_cap"))
    email_text = st.text_area(T("email_label"), height=220, placeholder=T("email_ph"))
    if st.button(T("email_btn"), type="primary", key="btn_email"):
        if not st.session_state.api_key:
            st.error(T("need_key"))
        elif not email_text.strip():
            st.warning(T("email_empty"))
        else:
            h = make_hash(email_text)
            if is_duplicate(h):
                st.warning(T("dup_warn"))
            else:
                with st.spinner(T("spin_email")):
                    try:
                        parsed = analyze_email(email_text)
                        if parsed.get("irrelevant"):
                            st.warning(T("irr_warn", reason=parsed.get("reason", T("irr_email"))))
                            st.session_state.input_hashes.append(h)
                        else:
                            new_b = parsed.get("new_banks", [])
                            apply_update(parsed, input_hash=h)
                            msg = T("update_ok")
                            if new_b: msg += T("new_banks", banks=', '.join(new_b))
                            st.success(msg); st.rerun()
                    except Exception as e:
                        st.error(T("err", e=e))

with inp2:
    st.caption(T("excel_cap"))
    uploaded_xlsx = st.file_uploader(T("excel_upload"), type=["xlsx"])
    if uploaded_xlsx:
        st.info(T("excel_uploaded", name=uploaded_xlsx.name))
        if st.button(T("excel_btn"), type="primary", key="btn_excel"):
            if not st.session_state.api_key:
                st.error(T("need_key"))
            else:
                file_bytes = uploaded_xlsx.read()
                h = make_hash(file_bytes.decode("latin-1"))
                if is_duplicate(h):
                    st.warning(T("dup_file"))
                else:
                    with st.spinner(T("spin_excel")):
                        try:
                            parsed = analyze_excel(read_excel_bytes(file_bytes))
                            if parsed.get("irrelevant"):
                                st.warning(T("irr_warn", reason=parsed.get("reason", T("irr_excel"))))
                                st.session_state.input_hashes.append(h)
                            else:
                                new_b = parsed.get("new_banks", [])
                                apply_update(parsed, input_hash=h)
                                msg = T("excel_ok", bank=parsed.get('submitting_bank','?'), cnt=parsed.get('bin_count',0))
                                if new_b: msg += T("new_banks", banks=', '.join(new_b))
                                st.success(msg); st.rerun()
                        except Exception as e:
                            st.error(T("err", e=e))

with inp3:
    st.caption(T("img_cap"))
    uploaded_img = st.file_uploader(T("img_upload"), type=["png","jpg","jpeg","webp"])
    if uploaded_img:
        st.image(uploaded_img, caption=f"📷 {uploaded_img.name}", use_container_width=True)
        if st.button(T("img_btn"), type="primary", key="btn_img"):
            if not st.session_state.api_key:
                st.error(T("need_key"))
            else:
                img_bytes  = uploaded_img.read()
                ext        = uploaded_img.name.rsplit(".", 1)[-1].lower()
                mt_map     = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","webp":"image/webp"}
                media_type = mt_map.get(ext, "image/png")
                h = make_hash(img_bytes.decode("latin-1"))
                if is_duplicate(h):
                    st.warning(T("dup_img"))
                else:
                    with st.spinner(T("spin_img")):
                        try:
                            parsed = analyze_image(img_bytes, media_type)
                            if parsed.get("irrelevant"):
                                st.warning(T("irr_warn", reason=parsed.get("reason", T("irr_img"))))
                                st.session_state.input_hashes.append(h)
                            else:
                                new_b = parsed.get("new_banks", [])
                                apply_update(parsed, input_hash=h)
                                msg = T("update_ok")
                                if new_b: msg += T("new_banks", banks=', '.join(new_b))
                                st.success(msg); st.rerun()
                        except Exception as e:
                            st.error(T("err", e=e))

st.divider()

# ── 히스토리 & 롤백 ────────────────────────────────────────────────────────────
if st.session_state.history:
    with st.expander(T("hist_exp", n=len(st.session_state.history)), expanded=False):
        st.caption(T("hist_cap"))
        for i, snap in reversed(list(enumerate(st.session_state.history))):
            ca, cb = st.columns([6, 1])
            ca.markdown(f'`{snap["time"]}` &nbsp; {snap["label"]} → **{snap["update_time"]}**')
            with cb:
                if st.button(T("restore_btn"), key=f"restore_{i}"):
                    restore_snapshot(i)
                    st.success(T("restore_ok"))
                    st.rerun()
    st.divider()

# ── Google Sheets 변경 이력 ────────────────────────────────────────────────────
st.subheader(T("trend_hdr"))
_gs2 = get_gsheet_client()
_sid = st.session_state.get("_sheet_id", "")

if not _gs2 or not _sid:
    st.info(T("trend_no_gs"))
else:
    with st.spinner(T("trend_spin")):
        hist_records = load_history_from_gsheet(_gs2, _sid)

    if not hist_records:
        st.info(T("trend_empty"))
    else:
        import pandas as pd
        _SC = {"완료":"#639922","진행중":"#EF9F27","확인중":"#185FA5","BIN 대기":"#EF9F27",
               "확인필요":"#EF9F27","미확인":"#888780","미완료":"#E24B4A","심각지연":"#E24B4A","신규":"#888780"}
        st.markdown(T("trend_total", n=len(hist_records)))
        df = pd.DataFrame(hist_records)
        if not df.empty:
            for _, row in df.iterrows():
                p_col = _SC.get(str(row.get("이전값","")), "#888780")
                c_col = _SC.get(str(row.get("현재값","")), "#639922")
                r1,r2,r3,r4,r5 = st.columns([2,1.5,1.5,1.5,1.5])
                r1.caption(str(row.get("날짜","")))
                r2.markdown(f"**{row.get('은행','')}**")
                r3.caption(str(row.get("항목","")))
                r4.markdown(f'<span style="background:{p_col};color:#fff;padding:1px 8px;border-radius:8px;font-size:11px;">{row.get("이전값","")}</span>', unsafe_allow_html=True)
                r5.markdown(f'<span style="background:{c_col};color:#fff;padding:1px 8px;border-radius:8px;font-size:11px;">→ {row.get("현재값","")}</span>', unsafe_allow_html=True)

st.divider()
st.caption(T("footer"))
